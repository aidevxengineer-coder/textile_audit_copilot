from __future__ import annotations

import asyncio
import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path
from typing import Any
from uuid import uuid4

import typer
from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from app.db.session import SessionLocal, create_db_and_seed
from app.models.upload import KnowledgeBaseDocument
from app.services.lightrag_service import LightRAGCompatibilityService
from app.services.ollama_service import OllamaService
from app.services.retrieval_service import RetrievalService
from app.services.schemas import GraphExtractionResult


app = typer.Typer(add_completion=False)

CHUNK_SIZE = 900
CHUNK_OVERLAP = 150
AUDIT_DATASETS = {
    "bangladesh_accord": Path("data/audit_test_data/bangladesh_accord"),
    "pakistan_accord": Path("data/audit_test_data/pakistan_accord"),
    "global_fla_brands": Path("data/audit_test_data/global_fla_brands"),
}
AUDIT_SPLIT_PATH = Path("data/processed/audit_dataset_split.json")


def read_manifest() -> list[dict[str, Any]]:
    return json.loads(Path("data/source_manifest.json").read_text(encoding="utf-8"))


def resolve_sources() -> list[dict[str, Any]]:
    resolved_manifest = Path("data/raw/downloads/resolved_manifest.json")
    if resolved_manifest.exists():
        return [
            source
            for source in json.loads(resolved_manifest.read_text(encoding="utf-8"))
            if source.get("local_path")
        ]
    sources = []
    for item in read_manifest():
        for extension in (".pdf", ".html", ".htm", ".txt", ".docx"):
            slug = re.sub(r"[^a-z0-9]+", "-", item["id"].lower()).strip("-")
            candidate = Path("data/raw/downloads") / f"{slug}{extension}"
            if candidate.exists():
                sources.append({**item, "local_path": str(candidate)})
                break
    return sources


def extract_text(file_path: Path) -> str:
    if file_path.suffix.lower() == ".pdf":
        reader = PdfReader(str(file_path))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        return "\n".join(pages)
    if file_path.suffix.lower() == ".docx":
        document = Document(str(file_path))
        return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    if file_path.suffix.lower() in {".html", ".htm"}:
        soup = BeautifulSoup(file_path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
        return soup.get_text("\n", strip=True)
    if file_path.suffix.lower() == ".xlsx":
        workbook = load_workbook(file_path, read_only=True, data_only=False, keep_links=False)
        output: list[str] = []
        try:
            for sheet in workbook.worksheets:
                output.append(f"## Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value).strip() for value in row[:200]]
                    while values and not values[-1]:
                        values.pop()
                    if values:
                        first = next((index for index, value in enumerate(values) if value), len(values))
                        values = values[first:]
                    if any(values):
                        output.append(" | ".join(values))
        finally:
            workbook.close()
        return "\n".join(output)
    return file_path.read_text(encoding="utf-8", errors="ignore")


def resolve_audit_sources(train_fraction: float = 0.8) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    train_fraction = min(max(train_fraction, 0.0), 1.0)
    train_sources: list[dict[str, Any]] = []
    split_manifest: dict[str, Any] = {"train_fraction": train_fraction, "datasets": {}}

    for dataset_name, root in AUDIT_DATASETS.items():
        manifest_path = root / "manifest.json"
        if not manifest_path.exists():
            continue
        rows = json.loads(manifest_path.read_text(encoding="utf-8"))
        usable_rows = [row for row in rows if row.get("local_path") and row.get("status") != "failed"]
        units: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in usable_rows:
            local_path = Path(row["local_path"])
            units[local_path.parent.name].append(row)

        ordered_units = sorted(
            units,
            key=lambda unit: hashlib.sha256(f"{dataset_name}:{unit}".encode()).hexdigest(),
        )
        test_count = 0 if not ordered_units or train_fraction >= 1 else max(1, round(len(ordered_units) * (1 - train_fraction)))
        test_units = set(ordered_units[:test_count])
        train_units = [unit for unit in ordered_units if unit not in test_units]
        split_manifest["datasets"][dataset_name] = {
            "train_units": train_units,
            "test_units": sorted(test_units),
            "train_documents": sum(len(units[unit]) for unit in train_units),
            "test_documents": sum(len(units[unit]) for unit in test_units),
        }

        for unit in train_units:
            for row in units[unit]:
                file_path = Path(row["local_path"])
                if not file_path.exists() or file_path.suffix.lower() not in {".pdf", ".xlsx"}:
                    continue
                document_type = row.get("report_type") or row.get("document_type") or file_path.stem
                title = row.get("summary") or row.get("factory") or unit.replace("_", " ")
                source_key = f"{dataset_name}:{unit}:{file_path.name}"
                train_sources.append(
                    {
                        "id": "audit-" + hashlib.sha256(source_key.encode()).hexdigest()[:20],
                        "url": row.get("source_url") or row.get("factory_page") or "",
                        "title": f"{title} — {document_type}",
                        "standard_name": f"Historical Audit Evidence ({dataset_name})",
                        "citation": f"{title}; {document_type}; public historical assessment",
                        "jurisdiction": "Pakistan" if dataset_name == "pakistan_accord" else "Bangladesh" if dataset_name == "bangladesh_accord" else "International",
                        "version": None,
                        "tags": ["historical-audit", dataset_name, str(document_type)],
                        "local_path": str(file_path),
                        "source_type": "historical_audit_example",
                        "authority_level": "example_not_standard",
                        "dataset": dataset_name,
                        "audit_unit": unit,
                        "document_type": str(document_type),
                        "split": "train",
                    }
                )

    AUDIT_SPLIT_PATH.parent.mkdir(parents=True, exist_ok=True)
    AUDIT_SPLIT_PATH.write_text(json.dumps(split_manifest, indent=2), encoding="utf-8")
    return train_sources, split_manifest


def chunk_text(text: str) -> list[str]:
    cleaned = re.sub(r"\s+", " ", text).strip()
    chunks: list[str] = []
    start = 0
    while start < len(cleaned):
        end = min(len(cleaned), start + CHUNK_SIZE)
        chunk = cleaned[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end == len(cleaned):
            break
        start = end - CHUNK_OVERLAP
    return chunks


def clear_existing_index(retrieval_service: RetrievalService) -> None:
    payload = retrieval_service.collection.get(include=[])
    ids = payload.get("ids", [])
    if ids:
        retrieval_service.collection.delete(ids=ids)


def fallback_graph_extraction(chunk_id: str, text: str, metadata: dict[str, Any]) -> GraphExtractionResult:
    text_lower = text.lower()
    keyword_specs = {
        "fire_safety": ("Fire Safety", "fire safety", ["extinguisher", "fire exit", "evacuation", "alarm", "blocked exit"]),
        "chemical_storage": ("Chemical Storage", "chemical storage", ["chemical", "msds", "spill", "solvent", "storage"]),
        "ppe": ("PPE", "personal protective equipment", ["ppe", "gloves", "mask", "helmet", "goggles"]),
        "machine_guarding": ("Machine Guarding", "machine safety", ["machine guard", "guarding", "needle guard", "lockout"]),
        "wages": ("Wages", "wages and compensation", ["wage", "overtime", "minimum wage", "salary"]),
        "working_hours": ("Working Hours", "working hours", ["working hours", "overtime", "rest day", "weekly hours"]),
        "child_labour": ("Child Labour", "child labour", ["child labour", "minimum age", "young worker"]),
        "forced_labour": ("Forced Labour", "forced labour", ["forced labour", "bonded labour", "retention of documents"]),
        "worker_records": ("Worker Records", "records and contracts", ["contract", "personnel file", "wage record", "attendance"]),
        "worker_welfare": ("Worker Welfare", "worker welfare", ["toilet", "potable water", "canteen", "dormitory"]),
        "freedom_of_association": ("Freedom of Association", "worker representation", ["union", "collective bargaining", "worker representative"]),
        "grievance_mechanism": ("Grievance Mechanism", "worker grievance", ["grievance", "complaint", "speak for change"]),
        "harmful_substances": ("Harmful Substances", "harmful substances", ["harmful substances", "substances list", "limit values"]),
    }

    standard_id = re.sub(r"[^a-z0-9]+", "_", metadata["standard_name"].lower()).strip("_")
    nodes = [
        {
            "id": standard_id,
            "name": metadata["standard_name"],
            "topic": metadata["standard_name"],
            "aliases": [metadata.get("citation") or metadata["standard_name"]],
            "standards": [metadata["standard_name"]],
            "chunk_ids": [chunk_id],
        }
    ]
    relationships = []

    for entity_id, (name, topic, hints) in keyword_specs.items():
        if any(hint in text_lower for hint in hints):
            nodes.append(
                {
                    "id": entity_id,
                    "name": name,
                    "topic": topic,
                    "aliases": hints[:2],
                    "standards": [metadata["standard_name"]],
                    "chunk_ids": [chunk_id],
                }
            )
            relationships.append(
                {
                    "source": standard_id,
                    "target": entity_id,
                    "type": "covers",
                    "weight": 1.0,
                    "evidence": metadata.get("citation") or metadata["title"],
                }
            )

    return GraphExtractionResult(entities=nodes, relationships=relationships)


async def build_graph(chunks: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    ollama_service = OllamaService()
    graph_service = LightRAGCompatibilityService(ollama_service=ollama_service)
    use_lightrag_core = graph_service.use_lightrag_core

    merged_nodes: dict[str, dict[str, Any]] = {}
    merged_edges: dict[tuple[str, str, str], dict[str, Any]] = {}
    fallback_count = 0

    for chunk in chunks:
        if use_lightrag_core:
            try:
                extraction = await graph_service.extract_graph_elements(chunk_id=chunk["id"], text=chunk["text"])
            except Exception:
                extraction = fallback_graph_extraction(chunk["id"], chunk["text"], chunk["metadata"])
                fallback_count += 1
        else:
            extraction = fallback_graph_extraction(chunk["id"], chunk["text"], chunk["metadata"])
            fallback_count += 1

        for entity in extraction.entities:
            entity_id = entity["id"]
            current = merged_nodes.setdefault(
                entity_id,
                {
                    "id": entity_id,
                    "name": entity.get("name", entity_id),
                    "topic": entity.get("topic", entity.get("name", entity_id)),
                    "aliases": [],
                    "standards": [],
                    "chunk_ids": [],
                },
            )
            current["aliases"] = sorted(set(current["aliases"]) | set(entity.get("aliases", [])))
            current["standards"] = sorted(set(current["standards"]) | set(entity.get("standards", [])))
            current["chunk_ids"] = sorted(set(current["chunk_ids"]) | set(entity.get("chunk_ids", [])))

        for relation in extraction.relationships:
            source = relation["source"]
            target = relation["target"]
            if source == target:
                continue
            key = tuple(sorted((source, target)) + [relation.get("type", "related")])
            current_edge = merged_edges.setdefault(
                key,
                {
                    "source": key[0],
                    "target": key[1],
                    "type": key[2],
                    "weight": 0.0,
                    "evidence": [],
                },
            )
            current_edge["weight"] += float(relation.get("weight", 1.0))
            evidence = relation.get("evidence")
            if evidence:
                current_edge["evidence"].append(evidence)

    nodes = list(merged_nodes.values())
    edges = []
    for edge in merged_edges.values():
        edge["evidence"] = sorted(set(edge["evidence"]))[:5]
        edges.append(edge)

    metadata = {
        "chunk_count": len(chunks),
        "node_count": len(nodes),
        "edge_count": len(edges),
        "fallback_chunks": fallback_count,
    }
    graph_service.write_lightrag_manifest(metadata)
    return nodes, edges, metadata


@app.command()
def ingest(
    skip_fetch: bool = typer.Option(False, "--skip-fetch", help="Assume data/raw/downloads is already populated."),
    include_audit_data: bool = typer.Option(
        False,
        "--include-audit-data",
        help="Add the training split of downloaded public audit examples to retrieval.",
    ),
    audit_train_fraction: float = typer.Option(0.8, "--audit-train-fraction", min=0.0, max=1.0),
    skip_db_sync: bool = typer.Option(
        False,
        "--skip-db-sync",
        help="Build retrieval and graph indexes without updating the PostgreSQL document catalog.",
    ),
    catalog_only: bool = typer.Option(
        False,
        "--catalog-only",
        help="Synchronize the PostgreSQL document catalog without rebuilding retrieval indexes.",
    ),
) -> None:
    async def _run() -> None:
        if catalog_only and skip_db_sync:
            raise typer.BadParameter("--catalog-only cannot be combined with --skip-db-sync")
        if not skip_db_sync:
            create_db_and_seed()

        if not skip_fetch:
            from scripts.fetch_compliance_docs import fetch

            fetch(force=False)

        official_sources = resolve_sources()
        audit_sources: list[dict[str, Any]] = []
        audit_split: dict[str, Any] | None = None
        if include_audit_data:
            audit_sources, audit_split = resolve_audit_sources(audit_train_fraction)
        sources = [*official_sources, *audit_sources]
        if not sources:
            raise typer.BadParameter("No source documents found. Run scripts/fetch_compliance_docs.py first.")

        if catalog_only:
            with SessionLocal() as db:
                db.query(KnowledgeBaseDocument).delete()
                for source in sources:
                    db.add(
                        KnowledgeBaseDocument(
                            source_url=source["url"],
                            title=source["title"],
                            local_path=str(source["local_path"]),
                            standard_name=source["standard_name"],
                            citation=source.get("citation"),
                            jurisdiction=source.get("jurisdiction"),
                            version=source.get("version"),
                            tags_json=source.get("tags", []),
                        )
                    )
                db.commit()
            typer.echo(json.dumps({"database_catalog_synced": True, "sources": len(sources)}, indent=2))
            return

        retrieval_service = RetrievalService()
        clear_existing_index(retrieval_service)

        chunks: list[dict[str, Any]] = []
        kb_records: list[KnowledgeBaseDocument] = []

        for source in sources:
            file_path = Path(source["local_path"])
            text = extract_text(file_path)
            if not text.strip():
                continue
            kb_records.append(
                KnowledgeBaseDocument(
                    source_url=source["url"],
                    title=source["title"],
                    local_path=str(file_path),
                    standard_name=source["standard_name"],
                    citation=source.get("citation"),
                    jurisdiction=source.get("jurisdiction"),
                    version=source.get("version"),
                    tags_json=source.get("tags", []),
                )
            )
            for index, chunk in enumerate(chunk_text(text)):
                chunk_id = str(uuid4())
                chunks.append(
                    {
                        "id": chunk_id,
                        "text": chunk,
                        "metadata": {
                            "source_id": source["id"],
                            "title": source["title"],
                            "standard_name": source["standard_name"],
                            "citation": source.get("citation") or "",
                            "clause_reference": f"{source['title']} - chunk {index + 1}",
                            "source_url": source["url"],
                            "version": source.get("version") or "",
                            "jurisdiction": source.get("jurisdiction") or "",
                            "source_type": source.get("source_type", "authoritative_standard"),
                            "authority_level": source.get("authority_level", "authoritative"),
                            "dataset": source.get("dataset", "official_corpus"),
                            "audit_unit": source.get("audit_unit", ""),
                            "document_type": source.get("document_type", "standard"),
                            "split": source.get("split", "train"),
                        },
                    }
                )

        retrieval_service.upsert_chunks(chunks)
        nodes, edges, graph_metadata = await build_graph(chunks)
        retrieval_service.save_graph(nodes=nodes, edges=edges)

        if not skip_db_sync:
            with SessionLocal() as db:
                db.query(KnowledgeBaseDocument).delete()
                for record in kb_records:
                    db.add(record)
                db.commit()

        summary = {
            "sources": len(kb_records),
            "official_sources": len(official_sources),
            "audit_training_sources": len(audit_sources),
            "database_catalog_synced": not skip_db_sync,
            "chunks": len(chunks),
            "graph": graph_metadata,
            "audit_split": audit_split,
        }
        Path("data/processed/ingest_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
        typer.echo(json.dumps(summary, indent=2))

    asyncio.run(_run())


if __name__ == "__main__":
    app()
