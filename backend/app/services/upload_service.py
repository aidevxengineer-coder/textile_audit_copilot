from __future__ import annotations

import csv
import hashlib
import io
import re
import shutil
import zipfile
from pathlib import Path
from typing import Any
from uuid import uuid4

import fitz
from PIL import Image
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy.orm import Session
from docx import Document

from app.config import get_settings
from app.core.encryption import decrypt_text, encrypt_text, write_encrypted_file
from app.models.upload import Upload
from app.services.schemas import AttachmentContext


ALLOWED_MIME_TYPES = {
    "image/jpeg",
    "image/jpg",
    "image/png",
    "image/webp",
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "text/plain",
    "text/csv",
    "application/csv",
}

DOCUMENT_TYPE_HINTS = {
    "policy": ["policy", "policies", "procedure", "sop", "handbook", "code of conduct"],
    "wages": ["wage", "salary", "payroll", "overtime", "minimum wage"],
    "attendance": ["attendance", "timesheet", "time sheet", "shift log"],
    "contracts": ["contract", "appointment", "employment letter", "agreement"],
    "safety": ["safety", "fire drill", "evacuation", "inspection", "certificate", "training"],
    "chemical": ["chemical", "msds", "sds", "hazmat", "inventory"],
    "grievance": ["grievance", "complaint", "disciplinary", "worker committee"],
    "welfare": ["canteen", "dormitory", "toilet", "drinking water", "welfare"],
}

PREVIEW_TEXT_LIMIT = 900
OCR_PAGE_LIMIT = 20
XLSX_MAX_SHEETS = 50
XLSX_MAX_ROWS_PER_SHEET = 10_000
XLSX_MAX_COLUMNS = 200
XLSX_MAX_NONEMPTY_CELLS = 100_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 5_000
MAX_PDF_PAGES = 1_000
MAX_IMAGE_PIXELS = 40_000_000
MAX_EMBEDDED_IMAGES_PER_DOCUMENT = 6
MIN_EMBEDDED_IMAGE_EDGE_PX = 120


class UploadValidationError(ValueError):
    pass


class UploadService:
    def __init__(self) -> None:
        self.settings = get_settings()

    def validate_upload(self, *, mime_type: str, size_bytes: int) -> None:
        if mime_type not in ALLOWED_MIME_TYPES:
            raise UploadValidationError("Only PNG, JPG, WEBP, PDF, DOCX, XLSX, CSV, and TXT files are allowed.")
        if size_bytes > self.settings.max_upload_mb * 1024 * 1024:
            raise UploadValidationError(f"File exceeds {self.settings.max_upload_mb}MB limit.")

    @staticmethod
    def normalize_mime_type(filename: str, mime_type: str) -> str:
        # Browsers occasionally submit image files as image/jpg or without a
        # content type. Normalize the common safe image extensions before the
        # allow-list check so factory photos are not rejected unnecessarily.
        value = (mime_type or "").lower().strip()
        suffix = Path(filename).suffix.lower()
        # Some browsers identify Office Open XML containers as generic ZIP
        # files. The extension is safe to use here because openpyxl validates
        # the workbook structure before any content is persisted.
        if suffix == ".xlsx" and value in {
            "",
            "application/octet-stream",
            "application/zip",
            "application/x-zip-compressed",
        }:
            return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if value == "image/jpg":
            return "image/jpeg"
        if value and value != "application/octet-stream":
            return value
        return {
            ".jpg": "image/jpeg",
            ".jpeg": "image/jpeg",
            ".png": "image/png",
            ".webp": "image/webp",
        }.get(suffix, value)

    def sanitize_image(self, raw_bytes: bytes, mime_type: str) -> tuple[bytes, str, tuple[int, int]]:
        try:
            image = Image.open(io.BytesIO(raw_bytes))
            if image.width * image.height > MAX_IMAGE_PIXELS:
                raise UploadValidationError("Image dimensions exceed the 40-megapixel safety limit.")
            image.load()
        except Exception as exc:
            raise UploadValidationError("The image could not be opened. Please upload a valid PNG, JPG, or WEBP file.") from exc

        dimensions = image.size
        output = io.BytesIO()
        if mime_type == "image/png":
            image.convert("RGBA" if image.mode == "RGBA" else "RGB").save(output, format="PNG", optimize=True)
            stored_mime_type = "image/png"
        elif mime_type == "image/webp":
            image.convert("RGB").save(output, format="WEBP", quality=90, method=6)
            stored_mime_type = "image/webp"
        else:
            image.convert("RGB").save(output, format="JPEG", optimize=True, quality=90)
            stored_mime_type = "image/jpeg"
        return output.getvalue(), stored_mime_type, dimensions

    @staticmethod
    def validate_office_archive(raw_bytes: bytes) -> None:
        try:
            with zipfile.ZipFile(io.BytesIO(raw_bytes)) as archive:
                entries = archive.infolist()
                if len(entries) > MAX_ARCHIVE_ENTRIES:
                    raise UploadValidationError("Office document contains too many archive entries.")
                total = sum(entry.file_size for entry in entries)
                if total > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                    raise UploadValidationError("Office document expands beyond the 100MB safety limit.")
                for entry in entries:
                    normalized = entry.filename.replace("\\", "/")
                    if normalized.startswith("/") or "../" in f"/{normalized}":
                        raise UploadValidationError("Office document contains an unsafe archive path.")
                    if entry.compress_size and entry.file_size / entry.compress_size > 200:
                        raise UploadValidationError("Office document has an unsafe compression ratio.")
        except zipfile.BadZipFile as exc:
            raise UploadValidationError("The Office document is not a valid archive.") from exc

    def extract_pdf_text(self, raw_bytes: bytes) -> str:
        try:
            with fitz.open(stream=raw_bytes, filetype="pdf") as document:
                parts = []
                for page_number, page in enumerate(document, start=1):
                    text = page.get_text("text").strip()
                    if text:
                        parts.append(f"[Page {page_number}]\n{text}")
                return "\n".join(parts).strip()
        except Exception as exc:
            raise UploadValidationError("The PDF could not be read. It may be damaged or password protected.") from exc

    def count_pdf_pages(self, raw_bytes: bytes) -> int:
        try:
            with fitz.open(stream=raw_bytes, filetype="pdf") as document:
                if document.page_count > MAX_PDF_PAGES:
                    raise UploadValidationError(f"PDF exceeds the {MAX_PDF_PAGES:,}-page safety limit.")
                return document.page_count
        except Exception:
            return 0

    def extract_embedded_images_from_pdf(
        self, raw_bytes: bytes
    ) -> list[tuple[bytes, str, int | None]]:
        """Pull raster images out of a PDF's pages.

        ``extract_pdf_text`` only reads the text layer, so a factory photo or a
        scanned certificate pasted into a document would otherwise never reach
        the vision model. Tiny images (icons, bullets, logos) are skipped since
        they add noise rather than compliance-relevant evidence.
        """

        found: list[tuple[bytes, str, int | None]] = []
        try:
            with fitz.open(stream=raw_bytes, filetype="pdf") as document:
                for page_number, page in enumerate(document, start=1):
                    if len(found) >= MAX_EMBEDDED_IMAGES_PER_DOCUMENT:
                        break
                    for image_ref in page.get_images(full=True):
                        if len(found) >= MAX_EMBEDDED_IMAGES_PER_DOCUMENT:
                            break
                        try:
                            extracted = document.extract_image(image_ref[0])
                            image_bytes = extracted["image"]
                            image = Image.open(io.BytesIO(image_bytes))
                            width, height = image.size
                            if min(width, height) < MIN_EMBEDDED_IMAGE_EDGE_PX:
                                continue
                            if width * height > MAX_IMAGE_PIXELS:
                                continue
                            ext = (extracted.get("ext") or "png").lower()
                            mime_type = "image/jpeg" if ext in {"jpg", "jpeg"} else f"image/{ext}"
                            found.append((image_bytes, mime_type, page_number))
                        except Exception:
                            continue
        except Exception:
            return []
        return found

    def extract_embedded_images_from_docx(
        self, raw_bytes: bytes
    ) -> list[tuple[bytes, str, int | None]]:
        """Pull raster images out of a DOCX package's ``word/media`` entries."""

        found: list[tuple[bytes, str, int | None]] = []
        try:
            with zipfile.ZipFile(io.BytesIO(raw_bytes)) as archive:
                media_names = sorted(
                    name
                    for name in archive.namelist()
                    if name.startswith("word/media/")
                    and Path(name).suffix.lower() in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
                )
                for name in media_names:
                    if len(found) >= MAX_EMBEDDED_IMAGES_PER_DOCUMENT:
                        break
                    try:
                        raw_image = archive.read(name)
                        image = Image.open(io.BytesIO(raw_image))
                        width, height = image.size
                        if min(width, height) < MIN_EMBEDDED_IMAGE_EDGE_PX:
                            continue
                        if width * height > MAX_IMAGE_PIXELS:
                            continue
                        suffix = Path(name).suffix.lower().lstrip(".")
                        mime_type = "image/jpeg" if suffix == "jpg" else f"image/{suffix}"
                        found.append((raw_image, mime_type, None))
                    except Exception:
                        continue
        except Exception:
            return []
        return found

    def _persist_embedded_images(
        self, *, images: list[tuple[bytes, str, int | None]], session_id: str
    ) -> list[dict[str, Any]]:
        records: list[dict[str, Any]] = []
        for index, (raw_image, mime_type, page_number) in enumerate(images, start=1):
            try:
                sanitized_bytes, stored_mime_type, dimensions = self.sanitize_image(raw_image, mime_type)
            except UploadValidationError:
                continue
            stored_path = Path(self.settings.upload_dir) / session_id / "embedded" / f"{uuid4()}.bin"
            write_encrypted_file(stored_path, sanitized_bytes)
            records.append(
                {
                    "stored_path": str(stored_path),
                    "mime_type": stored_mime_type,
                    "width": dimensions[0],
                    "height": dimensions[1],
                    "page": page_number,
                    "index": index,
                }
            )
        return records

    def try_ocr_pdf(self, raw_bytes: bytes) -> str | None:
        """Use OCR when it is already available on the host; otherwise fail closed.

        pytesseract and the native Tesseract executable are intentionally optional.
        A deployment without them still persists the encrypted PDF and exposes a
        clear ``needs_ocr`` state instead of claiming that extraction succeeded.
        """

        try:
            import pytesseract  # type: ignore[import-not-found]
        except ImportError:
            return None

        if not shutil.which(pytesseract.pytesseract.tesseract_cmd):
            # A package installer (e.g. winget) can put the binary on PATH for
            # new processes without this already-running process seeing the
            # refreshed PATH. Check the well-known install locations directly
            # rather than requiring a restart before OCR starts working.
            for candidate in (
                r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
                "/usr/bin/tesseract",
                "/usr/local/bin/tesseract",
            ):
                if Path(candidate).exists():
                    pytesseract.pytesseract.tesseract_cmd = candidate
                    break
            else:
                return None

        chunks: list[str] = []
        try:
            with fitz.open(stream=raw_bytes, filetype="pdf") as document:
                for page_number, page in enumerate(list(document)[:OCR_PAGE_LIMIT], start=1):
                    pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    image = Image.open(io.BytesIO(pixmap.tobytes("png")))
                    text = pytesseract.image_to_string(image).strip()
                    if text:
                        chunks.append(f"[Page {page_number} | OCR]\n{text}")
        except Exception:
            return None

        result = "\n".join(chunks).strip()
        return result or None

    def extract_docx_text(self, raw_bytes: bytes) -> str:
        try:
            document = Document(io.BytesIO(raw_bytes))
        except Exception as exc:
            raise UploadValidationError("The DOCX file could not be read. Please upload a valid Word document.") from exc
        parts = [
            f"[Paragraph {index}] {paragraph.text.strip()}"
            for index, paragraph in enumerate(document.paragraphs, start=1)
            if paragraph.text.strip()
        ]
        for table_number, table in enumerate(document.tables, start=1):
            for row_number, row in enumerate(table.rows, start=1):
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    parts.append(f"[Table {table_number} | Row {row_number}] " + " | ".join(values))
        return "\n".join(parts).strip()

    def extract_plain_text(self, raw_bytes: bytes) -> str:
        return raw_bytes.decode("utf-8", errors="ignore").strip()

    def extract_csv_text(self, raw_bytes: bytes) -> str:
        text = raw_bytes.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
        if len(rows) > 10000:
            raise UploadValidationError("CSV exceeds the 10,000 row safety limit.")
        return "\n".join(
            f"[Row {index}] " + " | ".join(cell.strip() for cell in row)
            for index, row in enumerate(rows, start=1)
        )

    def extract_xlsx_text(self, raw_bytes: bytes) -> tuple[str, dict[str, Any]]:
        try:
            workbook = load_workbook(
                io.BytesIO(raw_bytes),
                read_only=True,
                data_only=False,
                keep_links=False,
            )
        except (InvalidFileException, OSError, ValueError, KeyError) as exc:
            raise UploadValidationError(
                "The XLSX file could not be read. It may be damaged or password protected."
            ) from exc

        if len(workbook.worksheets) > XLSX_MAX_SHEETS:
            workbook.close()
            raise UploadValidationError(f"XLSX exceeds the {XLSX_MAX_SHEETS}-sheet safety limit.")

        output: list[str] = []
        nonempty_cells = 0
        rows_read = 0
        try:
            for sheet in workbook.worksheets:
                output.append(f"## Sheet: {sheet.title}")
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    if row_number > XLSX_MAX_ROWS_PER_SHEET:
                        raise UploadValidationError(
                            f"Sheet '{sheet.title}' exceeds the {XLSX_MAX_ROWS_PER_SHEET:,}-row safety limit."
                        )
                    values: list[str] = []
                    for value in row[:XLSX_MAX_COLUMNS]:
                        if value is None:
                            values.append("")
                            continue
                        nonempty_cells += 1
                        if nonempty_cells > XLSX_MAX_NONEMPTY_CELLS:
                            raise UploadValidationError(
                                f"XLSX exceeds the {XLSX_MAX_NONEMPTY_CELLS:,}-cell safety limit."
                            )
                        values.append(str(value).strip())
                    while values and not values[-1]:
                        values.pop()
                    if values:
                        first_value = next((index for index, value in enumerate(values) if value), len(values))
                        values = values[first_value:]
                    if any(values):
                        output.append(f"[Sheet: {sheet.title} | Row: {row_number}] " + " | ".join(values))
                        rows_read += 1
        finally:
            workbook.close()

        return "\n".join(output).strip(), {
            "sheet_count": len(workbook.worksheets),
            "sheet_names": [sheet.title for sheet in workbook.worksheets],
            "spreadsheet_rows_read": rows_read,
            "spreadsheet_nonempty_cells": nonempty_cells,
        }

    def infer_document_type(self, filename: str, extracted_text: str | None) -> str | None:
        haystack = f"{filename} {extracted_text or ''}".lower()
        for document_type, hints in DOCUMENT_TYPE_HINTS.items():
            if any(hint in haystack for hint in hints):
                return document_type
        return None

    def build_evidence_summary(
        self,
        *,
        filename: str,
        mime_type: str,
        document_type: str | None,
        extracted_text: str | None,
        embedded_image_count: int = 0,
    ) -> str:
        if mime_type.startswith("image/"):
            return f"Image evidence uploaded: {filename}."

        image_suffix = (
            f" ({embedded_image_count} embedded image(s) attached for visual review)"
            if embedded_image_count
            else ""
        )
        if extracted_text:
            preview = " ".join(extracted_text.split())[:320]
            kind = document_type or "document"
            return f"{kind.title()} evidence from {filename}: {preview}{image_suffix}"

        return f"Document uploaded: {filename}.{image_suffix}"

    @staticmethod
    def _preview_text(value: str | None, *, limit: int = PREVIEW_TEXT_LIMIT) -> str | None:
        if not value:
            return None
        normalized = " ".join(value.split())
        if len(normalized) <= limit:
            return normalized
        return f"{normalized[:limit].rstrip()}..."

    @staticmethod
    def read_extracted_text(upload: Upload) -> str | None:
        value = upload.extracted_text
        if not value:
            return None
        if (upload.metadata_json or {}).get("extracted_text_encrypted"):
            try:
                return decrypt_text(value)
            except Exception:
                return None
        return value

    @staticmethod
    def _effective_ingestion_status(upload: Upload) -> str:
        status = upload.ingestion_status or "ready"
        extracted_text = UploadService.read_extracted_text(upload)
        # Existing rows predate the explicit lifecycle fields. Infer an honest
        # state for them so the UI never presents an empty extraction as ready.
        if upload.mime_type.startswith("image/") and not extracted_text and status == "ready":
            return "ready_for_visual_review"
        if upload.mime_type == "application/pdf" and not (extracted_text or "").strip() and status == "ready":
            return "needs_ocr"
        return status

    def serialize_upload(self, upload: Upload) -> dict[str, Any]:
        metadata = upload.metadata_json or {}
        extracted_text = self.read_extracted_text(upload)
        ingestion_status = self._effective_ingestion_status(upload)
        preview_available = upload.mime_type.startswith("image/") or upload.mime_type in {
            "application/pdf",
            "text/plain",
            "text/csv",
            "application/csv",
        }
        extraction_method = upload.extraction_method
        if not extraction_method:
            if upload.mime_type.startswith("image/"):
                extraction_method = "vision"
            elif ingestion_status == "needs_ocr":
                extraction_method = "ocr"
            elif extracted_text:
                extraction_method = "native"

        extraction_message = upload.extraction_message
        if not extraction_message:
            if ingestion_status == "ready_for_visual_review":
                extraction_message = "Image attached successfully and ready for visual review."
            elif ingestion_status == "needs_ocr":
                extraction_message = "This PDF appears to be scanned and needs OCR before its text can be searched."
            elif ingestion_status == "failed":
                extraction_message = "The file was saved securely, but its contents could not be extracted."
            elif extracted_text:
                extraction_message = "Text extracted successfully and ready for project questions."
            else:
                extraction_message = "File attached successfully."

        return {
            "id": upload.id,
            "filename": upload.original_filename,
            "mime_type": upload.mime_type,
            "document_type": metadata.get("document_type"),
            "size_bytes": upload.size_bytes,
            "created_at": upload.created_at.isoformat() if upload.created_at else None,
            "has_extracted_text": bool((extracted_text or "").strip()),
            "preview_text": self._preview_text(extracted_text),
            "ingestion_status": ingestion_status,
            # Kept as an alias while the frontend moves to ingestion_status.
            "extraction_status": ingestion_status,
            "extraction_method": extraction_method,
            "extraction_message": extraction_message,
            "error_message": upload.error_message,
            "page_count": metadata.get("page_count"),
            "duplicate_of_upload_id": metadata.get("duplicate_of_upload_id"),
            "possible_version_of_upload_id": metadata.get("possible_version_of_upload_id"),
            "image_width": metadata.get("image_width"),
            "image_height": metadata.get("image_height"),
            "preview_available": preview_available,
            "preview_url": f"/api/uploads/{upload.id}/preview" if preview_available else None,
            "download_url": f"/api/uploads/{upload.id}/download",
        }

    def persist_upload(
        self,
        db: Session,
        *,
        owner_id: str,
        project_id: str | None,
        session_id: str,
        filename: str,
        mime_type: str,
        raw_bytes: bytes,
    ) -> Upload:
        mime_type = self.normalize_mime_type(filename, mime_type)
        self.validate_upload(mime_type=mime_type, size_bytes=len(raw_bytes))

        processed_bytes = raw_bytes
        extracted_text = None
        stored_mime_type = mime_type
        ingestion_status = "extracting"
        extraction_method: str | None = None
        extraction_message: str | None = None
        error_message: str | None = None
        metadata: dict[str, Any] = {"original_size_bytes": len(raw_bytes)}
        if mime_type.startswith("image/"):
            processed_bytes, stored_mime_type, dimensions = self.sanitize_image(raw_bytes, mime_type)
            metadata["sanitized_format"] = stored_mime_type.split("/", 1)[1]
            metadata["image_width"], metadata["image_height"] = dimensions
            ingestion_status = "ready_for_visual_review"
            extraction_method = "vision"
            extraction_message = "Image attached successfully and ready for visual review."
        elif mime_type == "application/pdf":
            extracted_text = self.extract_pdf_text(raw_bytes)
            metadata["page_count"] = self.count_pdf_pages(raw_bytes)
            embedded_images = self.extract_embedded_images_from_pdf(raw_bytes)
            if embedded_images:
                metadata["embedded_images"] = self._persist_embedded_images(
                    images=embedded_images, session_id=session_id
                )
            if extracted_text:
                ingestion_status = "ready"
                extraction_method = "native"
                extraction_message = "PDF text extracted successfully and ready for project questions."
            else:
                extracted_text = self.try_ocr_pdf(raw_bytes)
                extraction_method = "ocr"
                if extracted_text:
                    ingestion_status = "ready"
                    extraction_message = "Scanned PDF processed with OCR and ready for project questions."
                else:
                    ingestion_status = "needs_ocr"
                    extraction_message = "This PDF appears to be scanned and needs OCR before its text can be searched."
        elif mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            self.validate_office_archive(raw_bytes)
            extracted_text = self.extract_docx_text(raw_bytes)
            embedded_images = self.extract_embedded_images_from_docx(raw_bytes)
            if embedded_images:
                metadata["embedded_images"] = self._persist_embedded_images(
                    images=embedded_images, session_id=session_id
                )
            extraction_method = "native"
            if extracted_text:
                ingestion_status = "ready"
                extraction_message = "Word document text extracted successfully and ready for project questions."
            else:
                ingestion_status = "failed"
                error_message = "No readable text was found in this Word document."
                extraction_message = "The Word document was attached, but it did not contain readable text."
        elif mime_type in {"text/csv", "application/csv"}:
            extracted_text = self.extract_csv_text(raw_bytes)
            extraction_method = "native"
            if extracted_text:
                ingestion_status = "ready"
                extraction_message = "Spreadsheet rows extracted successfully and ready for project questions."
            else:
                ingestion_status = "failed"
                error_message = "No readable rows were found in this spreadsheet."
                extraction_message = "The spreadsheet was attached, but it did not contain readable rows."
        elif mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            self.validate_office_archive(raw_bytes)
            extracted_text, workbook_metadata = self.extract_xlsx_text(raw_bytes)
            metadata.update(workbook_metadata)
            extraction_method = "native"
            if extracted_text:
                ingestion_status = "ready"
                extraction_message = "Excel workbook rows extracted successfully and ready for project questions."
            else:
                ingestion_status = "failed"
                error_message = "No readable rows were found in this Excel workbook."
                extraction_message = "The Excel workbook was attached, but it did not contain readable rows."
        elif mime_type == "text/plain":
            extracted_text = self.extract_plain_text(raw_bytes)
            extraction_method = "native"
            if extracted_text:
                ingestion_status = "ready"
                extraction_message = "Text extracted successfully and ready for project questions."
            else:
                ingestion_status = "failed"
                error_message = "No readable text was found in this file."
                extraction_message = "The file was attached, but it did not contain readable text."

        if ingestion_status == "extracting":
            ingestion_status = "failed"
            error_message = "No extraction handler was available for this file."
            extraction_message = "The file was saved securely, but its contents could not be extracted."

        document_type = self.infer_document_type(filename, extracted_text)
        if document_type:
            metadata["document_type"] = document_type

        sha256 = hashlib.sha256(processed_bytes).hexdigest()
        prior_uploads = (
            db.query(Upload)
            .filter(Upload.owner_id == owner_id, Upload.project_id == project_id)
            .order_by(Upload.created_at.desc())
            .all()
        )
        duplicate = next((row for row in prior_uploads if row.sha256 == sha256), None)

        def version_key(value: str) -> str:
            stem = Path(value).stem.lower()
            stem = re.sub(r"(?:v(?:ersion)?|rev(?:ision)?)\s*[-_.]?\s*\d+(?:\.\d+)*", " ", stem)
            stem = re.sub(r"20\d{2}[-_.]?\d{0,2}[-_.]?\d{0,2}", " ", stem)
            return re.sub(r"[^a-z0-9]+", " ", stem).strip()

        current_key = version_key(filename)
        version_of = next(
            (row for row in prior_uploads if row.sha256 != sha256 and current_key and version_key(row.original_filename) == current_key),
            None,
        )
        if duplicate:
            metadata["duplicate_of_upload_id"] = duplicate.id
        elif version_of:
            metadata["possible_version_of_upload_id"] = version_of.id
        stored_extracted_text = encrypt_text(extracted_text) if extracted_text else None
        if stored_extracted_text:
            metadata["extracted_text_encrypted"] = True
        stored_name = f"{uuid4()}.bin"
        stored_path = Path(self.settings.upload_dir) / session_id / stored_name
        write_encrypted_file(stored_path, processed_bytes)

        upload = Upload(
            owner_id=owner_id,
            project_id=project_id,
            session_id=session_id,
            original_filename=filename,
            stored_path=str(stored_path),
            mime_type=stored_mime_type,
            size_bytes=len(processed_bytes),
            sha256=sha256,
            extracted_text=stored_extracted_text,
            ingestion_status=ingestion_status,
            extraction_method=extraction_method,
            extraction_message=extraction_message,
            error_message=error_message,
            metadata_json=metadata,
        )
        db.add(upload)
        db.commit()
        db.refresh(upload)
        return upload

    def to_context(self, upload: Upload) -> AttachmentContext:
        metadata = upload.metadata_json or {}
        embedded_images = metadata.get("embedded_images") or []
        extracted_text = self.read_extracted_text(upload)
        preview_text = self._preview_text(extracted_text, limit=1200)
        serialized = self.serialize_upload(upload)
        evidence_summary = self.build_evidence_summary(
            filename=upload.original_filename,
            mime_type=upload.mime_type,
            document_type=metadata.get("document_type"),
            extracted_text=extracted_text,
            embedded_image_count=len(embedded_images),
        )
        if serialized["ingestion_status"] == "needs_ocr":
            evidence_summary = f"{upload.original_filename}: {serialized['extraction_message']}"
        return AttachmentContext(
            upload_id=upload.id,
            file_name=upload.original_filename,
            mime_type=upload.mime_type,
            stored_path=upload.stored_path,
            document_type=metadata.get("document_type"),
            extracted_text=extracted_text,
            preview_text=preview_text,
            evidence_summary=evidence_summary,
            ingestion_status=serialized["ingestion_status"],
            extraction_method=serialized["extraction_method"],
            extraction_message=serialized["extraction_message"],
            embedded_images=embedded_images,
        )
