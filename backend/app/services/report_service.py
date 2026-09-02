from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.config import get_settings
from app.services.schemas import ComplianceReport


class ReportService:
    def __init__(self) -> None:
        self.settings = get_settings()

    def render_markdown(self, report: ComplianceReport) -> str:
        lines = [
            "# AuditReady AI Pre-Audit Report",
            "",
            f"> {report.disclaimer}",
            "",
            "## Executive Summary",
            report.executive_summary,
            "",
            "## Scope Reviewed",
            *[f"- {item}" for item in report.scope_reviewed],
            "",
            "## Scope Limitations",
            *[f"- {item}" for item in report.scope_limitations],
            "",
            "## Findings",
        ]
        for finding in report.findings:
            lines.extend(
                [
                    f"### {finding.category} - {finding.verdict}",
                    f"- Standard: {finding.standard_name}",
                    f"- Clause: {finding.clause_reference}",
                    f"- Confidence: {finding.confidence}",
                    f"- Priority: {finding.priority}",
                    f"- Audit Risk: {finding.audit_risk}",
                    f"- Evidence: {finding.evidence_summary}",
                    f"- Evidence Gap: {finding.evidence_gap}",
                    f"- Recommended Fix: {finding.remediation}",
                    *[f"- Expected Closure Evidence: {item}" for item in finding.recommended_evidence],
                    "",
                ]
            )
        lines.extend(["## Positive Controls", *[f"- {item}" for item in report.positive_controls], ""])
        lines.extend(["## Critical Missing Documents", *[f"- {item}" for item in report.critical_missing_documents], ""])
        lines.extend(["## Recommended Actions", *[f"- {item}" for item in report.recommended_actions], ""])
        lines.extend(["## Follow-up Questions", *[f"- {item}" for item in report.next_best_questions], ""])
        if report.guardrail_notes:
            lines.extend(["## Automated Evidence Checks", *[f"- {item}" for item in report.guardrail_notes], ""])
        return "\n".join(lines)

    def generate_pdf(self, report: ComplianceReport, report_id: str) -> Path:
        target = Path(self.settings.report_dir) / f"{report_id}.pdf"
        target.parent.mkdir(parents=True, exist_ok=True)
        doc = SimpleDocTemplate(str(target), pagesize=A4)
        styles = getSampleStyleSheet()
        body = ParagraphStyle("Body", parent=styles["BodyText"], leading=15)

        story = [
            Paragraph("AuditReady AI Pre-Audit Report", styles["Title"]),
            Spacer(1, 12),
            Paragraph(report.disclaimer, styles["Italic"]),
            Spacer(1, 12),
            Paragraph("Executive Summary", styles["Heading2"]),
            Paragraph(report.executive_summary, body),
            Spacer(1, 12),
            Paragraph("Findings", styles["Heading2"]),
        ]

        table_rows = [["Category", "Verdict", "Standard", "Clause", "Confidence"]]
        for finding in report.findings:
            table_rows.append(
                [
                    finding.category,
                    finding.verdict,
                    finding.standard_name,
                    finding.clause_reference,
                    finding.confidence,
                ]
            )
            story.append(Paragraph(f"{finding.category}: {finding.evidence_summary}", body))
            story.append(Paragraph(f"Recommended fix: {finding.remediation}", body))
            story.append(Spacer(1, 8))

        table = Table(table_rows, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#A91E22")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.lightgrey]),
                ]
            )
        )
        story.insert(6, table)
        story.extend(
            [
                Spacer(1, 12),
                Paragraph("Recommended Actions", styles["Heading2"]),
                *[Paragraph(f"- {item}", body) for item in report.recommended_actions],
            ]
        )
        doc.build(story)
        return target

    def generate_docx(self, report: ComplianceReport, report_id: str) -> Path:
        target = Path(self.settings.report_dir) / f"{report_id}.docx"
        target.parent.mkdir(parents=True, exist_ok=True)

        document = Document()
        document.add_heading("AuditReady AI Pre-Audit Report", level=0)
        document.add_paragraph(report.disclaimer)
        document.add_heading("Executive Summary", level=1)
        document.add_paragraph(report.executive_summary)
        document.add_heading("Findings", level=1)
        for finding in report.findings:
            document.add_heading(f"{finding.category}: {finding.verdict}", level=2)
            document.add_paragraph(f"Standard: {finding.standard_name}")
            document.add_paragraph(f"Clause: {finding.clause_reference}")
            document.add_paragraph(f"Confidence: {finding.confidence}")
            document.add_paragraph(f"Priority: {finding.priority}")
            document.add_paragraph(f"Audit risk: {finding.audit_risk}")
            document.add_paragraph(f"Evidence: {finding.evidence_summary}")
            document.add_paragraph(f"Evidence gap: {finding.evidence_gap}")
            document.add_paragraph(f"Recommended fix: {finding.remediation}")
            for evidence in finding.recommended_evidence:
                document.add_paragraph(f"Expected closure evidence: {evidence}")

        document.add_heading("Recommended Actions", level=1)
        for item in report.recommended_actions:
            document.add_paragraph(item, style="List Bullet")
        document.add_heading("Follow-up Questions", level=1)
        for item in report.next_best_questions:
            document.add_paragraph(item, style="List Bullet")
        document.save(target)
        return target

    def generate_csv(self, report: ComplianceReport, report_id: str) -> Path:
        target = Path(self.settings.report_dir) / f"{report_id}.csv"
        target.parent.mkdir(parents=True, exist_ok=True)
        with target.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=[
                    "category",
                    "verdict",
                    "standard",
                    "clause",
                    "confidence",
                    "priority",
                    "audit_risk",
                    "evidence",
                    "evidence_gap",
                    "recommended_evidence",
                    "recommended_fix",
                ],
            )
            writer.writeheader()
            for finding in report.findings:
                writer.writerow(
                    {
                        "category": finding.category,
                        "verdict": finding.verdict,
                        "standard": finding.standard_name,
                        "clause": finding.clause_reference,
                        "confidence": finding.confidence,
                        "priority": finding.priority,
                        "audit_risk": finding.audit_risk,
                        "evidence": finding.evidence_summary,
                        "evidence_gap": finding.evidence_gap,
                        "recommended_evidence": "; ".join(finding.recommended_evidence),
                        "recommended_fix": finding.remediation,
                    }
                )
        return target
