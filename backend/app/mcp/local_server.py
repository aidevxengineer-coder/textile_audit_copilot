from __future__ import annotations

import ast
import base64
import csv
import email
import imaplib
import io
import json
import operator
import os
import smtplib
import sys
import uuid
from datetime import datetime, timezone
from email.message import EmailMessage
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlparse

import httpx
from bs4 import BeautifulSoup
from docx import Document
from mcp.server.fastmcp import FastMCP
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from dotenv import load_dotenv

load_dotenv(override=False)

MODE = sys.argv[1] if len(sys.argv) > 1 else "utilities"
mcp = FastMCP(f"auditready-{MODE}")
ROOT = Path(os.getenv("MCP_WORKSPACE_ROOT", Path.cwd())).resolve()
OUTPUT = Path(os.getenv("MCP_OUTPUT_DIR", ROOT / "backend" / "generated")).resolve()
OUTPUT.mkdir(parents=True, exist_ok=True)

DUMMY_GMAIL_MESSAGES = [
    {
        "id": "gmail-001",
        "thread_id": "thread-audit-001",
        "from": "buyer.compliance@example.test",
        "to": "factory@example.test",
        "subject": "Upcoming social compliance audit",
        "body": "Please share the latest safety training register before the audit.",
        "received_at": "2026-07-10T09:30:00Z",
        "attachments": ["audit-request.txt"],
    },
    {
        "id": "gmail-002",
        "thread_id": "thread-audit-002",
        "from": "hr@example.test",
        "to": "compliance@example.test",
        "subject": "Training register updated",
        "body": "The July fire-safety attendance record is attached.",
        "received_at": "2026-07-12T12:15:00Z",
        "attachments": ["training-register.csv"],
    },
]

DUMMY_ATTACHMENTS = {
    ("gmail-001", "audit-request.txt"): b"Dummy audit request: provide the safety training register.\n",
    ("gmail-002", "training-register.csv"): b"employee,training,date\nAyesha,Fire Safety,2026-07-05\nBilal,Fire Safety,2026-07-05\n",
}

DUMMY_WHATSAPP_MESSAGES = [
    {
        "id": "wa-001",
        "chat_id": "factory-safety-team",
        "from": "+92-300-0000001",
        "sender_name": "Safety Officer",
        "body": "Fire exit inspection completed for Building A.",
        "sent_at": "2026-07-13T08:20:00Z",
        "direction": "inbound",
    },
    {
        "id": "wa-002",
        "chat_id": "factory-safety-team",
        "from": "+92-300-0000002",
        "sender_name": "Floor Supervisor",
        "body": "Two emergency lights need replacement on floor 2.",
        "sent_at": "2026-07-13T08:27:00Z",
        "direction": "inbound",
    },
]


def append_dummy_event(filename: str, payload: dict) -> Path:
    target = OUTPUT / filename
    with target.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(payload, ensure_ascii=False) + "\n")
    return target


def safe_path(value: str, *, output: bool = False) -> Path:
    base = OUTPUT if output else ROOT
    candidate = Path(value)
    candidate = (base / candidate).resolve() if not candidate.is_absolute() else candidate.resolve()
    allowed = (base, OUTPUT)
    if not any(candidate == root or root in candidate.parents for root in allowed):
        raise ValueError("Path is outside the allowed workspace.")
    return candidate


@mcp.tool()
def health() -> dict:
    """Return server health and capability mode."""
    return {"ok": True, "server": MODE, "open_source": True}


if MODE == "filesystem":
    @mcp.tool()
    def write_file(path: str, content_base64: str) -> dict:
        """Write a base64 file inside the workspace/output directory."""
        target = safe_path(path)
        target.parent.mkdir(parents=True, exist_ok=True)
        payload = base64.b64decode(content_base64, validate=True)
        target.write_bytes(payload)
        return {"path": str(target), "size_bytes": len(payload)}

    @mcp.tool()
    def read_file(path: str) -> dict:
        """Read a workspace file as base64."""
        target = safe_path(path)
        payload = target.read_bytes()
        return {"path": str(target), "content_base64": base64.b64encode(payload).decode(), "size_bytes": len(payload)}


if MODE == "documents":
    @mcp.tool()
    def generate_document(format: str, title: str, content: str, filename: str) -> dict:
        """Generate PDF, DOCX, CSV, TXT, or Markdown output."""
        kind = format.lower().lstrip(".")
        if kind not in {"pdf", "docx", "csv", "txt", "md"}:
            raise ValueError("Supported formats: pdf, docx, csv, txt, md")
        target = safe_path(str(Path(filename).with_suffix(f".{kind}")), output=True)
        target.parent.mkdir(parents=True, exist_ok=True)
        if kind == "docx":
            doc = Document(); doc.add_heading(title, 0)
            for paragraph in content.splitlines(): doc.add_paragraph(paragraph)
            doc.save(target)
        elif kind == "pdf":
            pdf = canvas.Canvas(str(target), pagesize=A4); width, height = A4; y = height - 60
            pdf.setFont("Helvetica-Bold", 16); pdf.drawString(50, y, title[:90]); y -= 30
            pdf.setFont("Helvetica", 10)
            for line in content.splitlines() or [""]:
                for start in range(0, max(1, len(line)), 100):
                    if y < 50: pdf.showPage(); pdf.setFont("Helvetica", 10); y = height - 50
                    pdf.drawString(50, y, line[start:start + 100]); y -= 14
            pdf.save()
        elif kind == "csv":
            rows = list(csv.reader(io.StringIO(content)))
            with target.open("w", newline="", encoding="utf-8-sig") as handle: csv.writer(handle).writerows(rows)
        else:
            target.write_text((f"# {title}\n\n" if kind == "md" else f"{title}\n\n") + content, encoding="utf-8")
        return {"path": str(target), "format": kind, "size_bytes": target.stat().st_size}


if MODE == "web_search":
    @mcp.tool()
    async def web_search(query: str, limit: int = 5) -> dict:
        """Search the public web through a configured SearXNG instance or DuckDuckGo HTML."""
        searxng = os.getenv("SEARXNG_URL", "").rstrip("/")
        try:
            provider_timeout = float(os.getenv("WEB_SEARCH_HTTP_TIMEOUT_SECONDS", "2.5"))
        except ValueError:
            provider_timeout = 2.5
        provider_timeout = max(1.0, min(provider_timeout, 10.0))
        async with httpx.AsyncClient(
            timeout=httpx.Timeout(provider_timeout, connect=min(1.5, provider_timeout)),
            follow_redirects=True,
            headers={
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126 Safari/537.36"
                )
            },
        ) as client:
            if searxng:
                response = await client.get(f"{searxng}/search", params={"q": query, "format": "json"})
                response.raise_for_status()
                rows = response.json().get("results", [])[: max(1, min(limit, 10))]
                return {"results": [{"title": r.get("title"), "url": r.get("url"), "snippet": r.get("content")} for r in rows]}
            response = await client.get("https://html.duckduckgo.com/html/", params={"q": query})
            response.raise_for_status()
        soup = BeautifulSoup(response.text, "html.parser")
        results = []
        for node in soup.select(".result")[: max(1, min(limit, 10))]:
            anchor = node.select_one(".result__a")
            if not anchor: continue
            href = anchor.get("href", "")
            if "uddg=" in href: href = unquote(parse_qs(urlparse(href).query).get("uddg", [href])[0])
            snippet = node.select_one(".result__snippet")
            results.append({"title": anchor.get_text(" ", strip=True), "url": href, "snippet": snippet.get_text(" ", strip=True) if snippet else ""})
        return {"results": results}


if MODE == "email":
    def build_message(recipient: str, subject: str, body: str, attachments: list[str]) -> EmailMessage:
        msg = EmailMessage(); msg["To"] = recipient; msg["From"] = os.getenv("SMTP_FROM", os.getenv("SMTP_USERNAME", "")); msg["Subject"] = subject; msg.set_content(body)
        for value in attachments:
            path = safe_path(value); msg.add_attachment(path.read_bytes(), maintype="application", subtype="octet-stream", filename=path.name)
        return msg

    @mcp.tool()
    def draft_email(recipient: str, subject: str, body: str, attachments: list[str] = []) -> dict:
        """Create a local RFC822 .eml draft without sending it."""
        msg = build_message(recipient, subject, body, attachments)
        target = OUTPUT / f"draft-{abs(hash((recipient, subject, body)))}.eml"; target.write_bytes(msg.as_bytes())
        return {"status": "drafted", "path": str(target)}

    @mcp.tool()
    def send_email(recipient: str, subject: str, body: str, attachments: list[str] = []) -> dict:
        """Simulate sending a Gmail message and store it in a local dummy outbox."""
        event = {"id": f"gmail-out-{uuid.uuid4().hex[:12]}", "recipient": recipient, "subject": subject,
                 "body": body, "attachments": attachments, "sent_at": datetime.now(timezone.utc).isoformat(),
                 "dummy": True}
        path = append_dummy_event("dummy-gmail-outbox.jsonl", event)
        return {"status": "simulated", "provider": "gmail", "dummy": True, "message": event, "outbox_path": str(path)}

    @mcp.tool()
    def list_messages(limit: int = 10) -> dict:
        """List deterministic dummy Gmail messages."""
        return {"provider": "gmail", "dummy": True, "messages": DUMMY_GMAIL_MESSAGES[:max(1, min(limit, 50))]}

    @mcp.tool()
    def download_attachment(message_id: str, attachment_name: str) -> dict:
        """Materialize one attachment from the dummy Gmail inbox."""
        payload = DUMMY_ATTACHMENTS.get((message_id, attachment_name))
        if payload is None:
            raise ValueError("Dummy attachment not found.")
        target = safe_path(Path(attachment_name).name, output=True); target.write_bytes(payload)
        return {"path": str(target), "size_bytes": target.stat().st_size, "dummy": True}


if MODE == "whatsapp":
    @mcp.tool()
    def list_messages(chat_id: str | None = None, limit: int = 20) -> dict:
        """List deterministic dummy WhatsApp messages, optionally filtered by chat."""
        rows = [row for row in DUMMY_WHATSAPP_MESSAGES if not chat_id or row["chat_id"] == chat_id]
        return {"provider": "whatsapp", "dummy": True, "messages": rows[:max(1, min(limit, 50))]}

    @mcp.tool()
    def send_message(recipient: str, body: str) -> dict:
        """Simulate a WhatsApp message and store it in a local dummy outbox."""
        event = {"id": f"wa-out-{uuid.uuid4().hex[:12]}", "recipient": recipient, "body": body,
                 "sent_at": datetime.now(timezone.utc).isoformat(), "direction": "outbound", "dummy": True}
        path = append_dummy_event("dummy-whatsapp-outbox.jsonl", event)
        return {"status": "simulated", "provider": "whatsapp", "dummy": True, "message": event, "outbox_path": str(path)}


if MODE == "utilities":
    OPS = {ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul, ast.Div: operator.truediv, ast.Pow: operator.pow, ast.USub: operator.neg}
    def evaluate(node):
        if isinstance(node, ast.Expression): return evaluate(node.body)
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)): return node.value
        if isinstance(node, ast.BinOp) and type(node.op) in OPS: return OPS[type(node.op)](evaluate(node.left), evaluate(node.right))
        if isinstance(node, ast.UnaryOp) and type(node.op) in OPS: return OPS[type(node.op)](evaluate(node.operand))
        raise ValueError("Unsupported expression")

    @mcp.tool()
    def calculator(query: str) -> dict:
        """Safely evaluate an arithmetic expression."""
        return {"result": evaluate(ast.parse(query, mode="eval"))}

    @mcp.tool()
    async def weather(query: str) -> dict:
        """Get live weather for a location using Open-Meteo."""
        async with httpx.AsyncClient(timeout=15) as client:
            geo = (await client.get("https://geocoding-api.open-meteo.com/v1/search", params={"name": query, "count": 1})).json().get("results", [])
            if not geo: raise ValueError("Location not found")
            place = geo[0]; current = (await client.get("https://api.open-meteo.com/v1/forecast", params={"latitude": place["latitude"], "longitude": place["longitude"], "current": "temperature_2m,relative_humidity_2m,wind_speed_10m"})).json().get("current", {})
        return {"location": f'{place["name"]}, {place.get("country", "")}', **current}


if __name__ == "__main__":
    mcp.run(transport="stdio")
